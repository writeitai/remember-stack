"""The D138 ``office`` route: Word documents and presentations → prose.

Word documents (docx, docm, dotx) are rendered with markitdown.
Presentations (pptx, pptm, ppsx, potx) are read slide by slide with
python-pptx: each slide becomes a ``## Slide N`` section with its text,
tables and speaker notes, and gets its own ``page`` locator (page = slide
number). The legacy and OpenDocument formats (doc, odt, rtf; ppt, odp) are
first converted to docx or pptx by LibreOffice. General D134 metadata (title,
creator, created, modified) comes from the package's ``docProps/core.xml``,
parsed with defusedxml. Everything the route writes is prose and is
claim-extracted; a file that is not a readable package fails the version.
"""

from datetime import datetime
from datetime import timezone
from importlib.metadata import version as package_version
import io
import logging
from typing import Final
from xml.etree.ElementTree import Element
from xml.etree.ElementTree import ParseError
import zipfile
import zlib

from defusedxml import DefusedXmlException
from defusedxml import ElementTree
from lxml.etree import XMLSyntaxError
from markitdown import MarkItDown
from markitdown import StreamInfo
from markitdown._exceptions import MarkItDownException
from markitdown.converters import DocxConverter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PythonPptxError
from pptx.shapes.base import BaseShape
from pptx.shapes.shapetree import SlideShapes
from pptx.slide import Slide

from rememberstack.adapters.converters.libreoffice import convert_with_libreoffice
from rememberstack.adapters.converters.time_limit import run_with_time_limit
from rememberstack.adapters.converters.zip_budget import require_zip_within_budget
from rememberstack.core import entire_document_labeling
from rememberstack.core.format_registry import family_for_mime
from rememberstack.core.format_registry import family_named
from rememberstack.core.format_registry import FORMAT_MIMES
from rememberstack.model import ConversionCoverage
from rememberstack.model import ConversionError
from rememberstack.model import ConversionResult
from rememberstack.model import ConverterManifest
from rememberstack.model import ManifestComponent
from rememberstack.model import PageLocator
from rememberstack.model import SourceMapEntry
from rememberstack.model.document_metadata import DocumentMetadata
from rememberstack.model.document_metadata import DocumentPerson

_logger = logging.getLogger(__name__)

OFFICE_CONVERTER_VERSION: Final = "office-2026.09"
"""Pins the office route: markitdown Word rendering, the per-slide
presentation shape, core-properties metadata and LibreOffice pre-conversion."""

_DOCX: Final = family_named(name="word").mime
_LEGACY_EXTENSIONS: Final = {
    FORMAT_MIMES[extension]: extension
    for extension in ("doc", "odt", "rtf", "ppt", "odp")
}
"""Stored MIME → source extension of the formats LibreOffice converts first."""

_PRESENTATION_MAIN: Final = (
    b"application/vnd.openxmlformats-officedocument.presentationml.presentation"
    b".main+xml"
)
_OTHER_PRESENTATION_MAINS: Final = (
    b"application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
)
"""Slide shows (ppsx) and templates (potx) differ from a pptx only in the
main part's content type, which python-pptx refuses; it is rewritten."""

_CORE_NAMESPACES: Final = {
    "dc": "http://purl.org/dc/elements/1.1/",
    "dcterms": "http://purl.org/dc/terms/",
}


class OfficeConverter:
    """Read Word documents and presentations, converting legacy files first."""

    def __init__(self) -> None:
        """Build a markitdown instance that knows only Word documents.

        Without other converters, markitdown cannot silently read a broken
        package as a zip listing or plain text; it fails instead.
        """
        self._markitdown = MarkItDown(enable_builtins=False, enable_plugins=False)
        self._markitdown.register_converter(DocxConverter())

    @property
    def name(self) -> str:
        """The route name recorded on representations."""
        return "office"

    @property
    def version(self) -> str:
        """The pinned office route version (D38)."""
        return OFFICE_CONVERTER_VERSION

    def convert(self, *, content: bytes, mime: str) -> ConversionResult:
        """Convert one Word document or presentation to Markdown."""
        family = family_for_mime(mime=mime).name
        if family not in ("word", "presentation"):
            raise ConversionError(f"the office route does not read {mime!r}")
        components = [_component(name="office", version=OFFICE_CONVERTER_VERSION)]
        legacy_extension = _LEGACY_EXTENSIONS.get(mime)
        if legacy_extension is not None:
            content = convert_with_libreoffice(
                content=content,
                source_extension=legacy_extension,
                target="docx" if family == "word" else "pptx",
            )
            components.append(_component(name="libreoffice", version="system"))
        if family == "word":
            return run_with_time_limit(
                work=lambda: self._word(content=content, components=components),
                what="Word document conversion",
            )
        return run_with_time_limit(
            work=lambda: _presentation(content=content, components=components),
            what="presentation conversion",
        )

    def _word(
        self, *, content: bytes, components: list[ManifestComponent]
    ) -> ConversionResult:
        """Render a Word document with markitdown."""
        warnings: list[str] = []
        require_zip_within_budget(content=content, what="Word document")
        metadata = _core_metadata(content=content, warnings=warnings)
        try:
            result = self._markitdown.convert_stream(
                io.BytesIO(content),
                stream_info=StreamInfo(mimetype=_DOCX, extension=".docx"),
            )
        except (MarkItDownException, zipfile.BadZipFile, zlib.error) as err:
            raise ConversionError("the Word document could not be read") from err
        document_md = result.text_content
        components.append(
            _component(name="markitdown", version=package_version("markitdown"))
        )
        return ConversionResult(
            document_md=document_md,
            metadata=metadata,
            warnings=tuple(warnings),
            manifest=ConverterManifest(
                components=tuple(components),
                coverage=ConversionCoverage(policy="office-text", complete=True),
                derivation_ranges=entire_document_labeling(
                    document_md=document_md,
                    derivation_kind="prose",
                    evidence_mode="source_expression",
                ),
            ),
        )


def _presentation(
    *, content: bytes, components: list[ManifestComponent]
) -> ConversionResult:
    """Read a presentation slide by slide; each slide is one page locator."""
    warnings: list[str] = []
    require_zip_within_budget(content=content, what="presentation")
    metadata = _core_metadata(content=content, warnings=warnings)
    components.append(
        _component(name="python-pptx", version=package_version("python-pptx"))
    )
    try:
        presentation = Presentation(
            io.BytesIO(_as_presentation_package(content=content))
        )
        sections = [
            _slide_markdown(number=number, slide=slide)
            for number, slide in enumerate(presentation.slides, start=1)
        ]
    except (
        PythonPptxError,
        KeyError,
        ValueError,
        XMLSyntaxError,
        zipfile.BadZipFile,
        zlib.error,
    ) as err:
        raise ConversionError("the presentation could not be read") from err
    document_md = ""
    source_map: list[SourceMapEntry] = []
    for number, section in enumerate(sections, start=1):
        start = len(document_md)
        document_md += section + "\n\n"
        source_map.append(
            SourceMapEntry(
                start=start,
                end=start + len(section),
                locators=(PageLocator(page=number, precision="page"),),
                region_kind="slide",
            )
        )
    return ConversionResult(
        document_md=document_md,
        metadata=metadata,
        warnings=tuple(warnings),
        source_map=tuple(source_map),
        manifest=ConverterManifest(
            components=tuple(components),
            coverage=ConversionCoverage(policy="office-slides", complete=True),
            derivation_ranges=entire_document_labeling(
                document_md=document_md,
                derivation_kind="prose",
                evidence_mode="source_expression",
            ),
        ),
    )


def _slide_markdown(*, number: int, slide: Slide) -> str:
    """One slide as a ``## Slide N: title`` section with its text and notes."""
    title_shape = slide.shapes.title
    title = (
        " ".join(title_shape.text_frame.text.split()) if title_shape is not None else ""
    )
    title_id = title_shape.shape_id if title_shape is not None else None
    parts = [f"## Slide {number}" + (f": {title}" if title else "")]
    parts.extend(
        text
        for shape in _all_shapes(shapes=slide.shapes)
        if shape.shape_id != title_id and (text := _shape_text(shape=shape))
    )
    if slide.has_notes_slide:
        notes_frame = slide.notes_slide.notes_text_frame
        notes = _clean(text=notes_frame.text) if notes_frame is not None else ""
        if notes:
            parts.append(f"Speaker notes:\n\n{notes}")
    return "\n\n".join(parts)


def _all_shapes(*, shapes: SlideShapes) -> list[BaseShape]:
    """Every shape on the slide in order, with group shapes flattened."""
    flattened: list[BaseShape] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flattened.extend(_all_shapes(shapes=shape.shapes))  # pyright: ignore[reportAttributeAccessIssue]
        else:
            flattened.append(shape)
    return flattened


def _shape_text(*, shape: BaseShape) -> str:
    """A shape's text, or its table as a Markdown table; empty when neither."""
    if shape.has_text_frame:
        return _clean(text=shape.text_frame.text)  # pyright: ignore[reportAttributeAccessIssue]
    if shape.has_table:
        rows = [
            [
                _clean(text=cell.text).replace("\n", " ").replace("|", "\\|")
                for cell in row.cells
            ]
            for row in shape.table.rows  # pyright: ignore[reportAttributeAccessIssue]
        ]
        if not rows:
            return ""
        lines = [f"| {' | '.join(rows[0])} |", f"|{' --- |' * len(rows[0])}"]
        lines.extend(f"| {' | '.join(row)} |" for row in rows[1:])
        return "\n".join(lines)
    return ""


def _clean(*, text: str) -> str:
    """Normalize python-pptx line breaks (vertical tab) and trim."""
    return text.replace("\v", "\n").strip()


def _as_presentation_package(*, content: bytes) -> bytes:
    """Rewrite a slide show's or template's main content type to a pptx's."""
    with zipfile.ZipFile(io.BytesIO(content)) as source:
        content_types = source.read("[Content_Types].xml")
        if not any(main in content_types for main in _OTHER_PRESENTATION_MAINS):
            return content
        rewritten = io.BytesIO()
        with zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info)
                if info.filename == "[Content_Types].xml":
                    for main in _OTHER_PRESENTATION_MAINS:
                        data = data.replace(main, _PRESENTATION_MAIN)
                target.writestr(info, data)
    return rewritten.getvalue()


def _core_metadata(*, content: bytes, warnings: list[str]) -> DocumentMetadata | None:
    """D134 metadata from ``docProps/core.xml``; a non-zip input is corrupt.

    Unreadable core properties lose only the metadata: a warning is recorded
    and the document is still read.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as package:
            if "docProps/core.xml" not in package.namelist():
                return None
            core_xml = package.read("docProps/core.xml")
    except (zipfile.BadZipFile, zlib.error) as err:
        raise ConversionError("the Office package could not be read") from err
    try:
        root = ElementTree.fromstring(core_xml)
    except (ParseError, DefusedXmlException) as err:
        _logger.warning("office core properties could not be parsed", exc_info=True)
        warnings.append(f"core properties not read ({type(err).__name__})")
        return None
    title = _core_text(root=root, tag="dc:title")
    creator = _core_text(root=root, tag="dc:creator")
    metadata = DocumentMetadata(
        title=title,
        authors=(DocumentPerson(name=creator),) if creator else (),
        created_at=_core_date(root=root, tag="dcterms:created"),
        modified_at=_core_date(root=root, tag="dcterms:modified"),
    )
    return metadata if metadata != DocumentMetadata() else None


def _core_text(*, root: Element, tag: str) -> str | None:
    """One core property's trimmed text, or None when absent or blank."""
    element = root.find(tag, _CORE_NAMESPACES)
    text = (element.text or "").strip() if element is not None else ""
    return text or None


def _core_date(*, root: Element, tag: str) -> datetime | None:
    """A W3CDTF core date in UTC; a date without a zone is read as UTC."""
    text = _core_text(root=root, tag=tag)
    if text is None:
        return None
    try:
        value = datetime.fromisoformat(text)
    except ValueError:
        _logger.warning("office core date %s is not ISO 8601: %r", tag, text)
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _component(*, name: str, version: str) -> ManifestComponent:
    """One local stage of the route's component graph."""
    return ManifestComponent(name=name, version=version, execution="library-local")
