"""D138 §5.1 data-file profiles: spreadsheets, delimited files and datasets.

Fixtures are generated here (openpyxl workbooks, CSV bytes, pyarrow and
pyreadstat files, SQLite databases), except one committed legacy ``.xls``,
which no dependency of the engine can write.
"""

from collections.abc import Callable
import dataclasses
import datetime
import io
from pathlib import Path
import re
import shutil
import sqlite3
import zipfile

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
import pandas as pd
import pyarrow as pa
import pyarrow.feather as feather
import pyarrow.parquet as pq
import pyreadstat
import pytest

from rememberstack.adapters.converters import build_conversion_routes
from rememberstack.adapters.converters import dataset as dataset_module
from rememberstack.adapters.converters import spreadsheet as spreadsheet_module
from rememberstack.adapters.converters import table as table_module
from rememberstack.adapters.converters.card import CardConverter
from rememberstack.adapters.converters.dataset import DatasetConverter
from rememberstack.adapters.converters.profile import infer_type
from rememberstack.adapters.converters.spreadsheet import SpreadsheetConverter
from rememberstack.adapters.converters.table import TableConverter
from rememberstack.core import blockize
from rememberstack.core import STOCK_CONVERSION_ROUTE_NAMES
from rememberstack.core.extraction_eligibility import block_eligibility
from rememberstack.core.extraction_eligibility import is_model_free
from rememberstack.core.format_registry import detect_mime
from rememberstack.core.format_registry import exceeds_reading_limit
from rememberstack.core.format_registry import family_for_mime
from rememberstack.core.format_registry import stock_route_names
from rememberstack.model import ConversionError
from rememberstack.model import ConversionResult
from rememberstack.model import FileHints

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLS = "application/vnd.ms-excel"
_ODS = "application/vnd.oasis.opendocument.spreadsheet"
_PARQUET = "application/vnd.apache.parquet"
_ARROW = "application/vnd.apache.arrow.file"
_SQLITE = "application/vnd.sqlite3"
_FIXTURES = Path(__file__).parent / "fixtures"


def _hints(name: str) -> FileHints:
    return FileHints(file_name=name, source_path=f"data/{name}")


def _xlsx(build: Callable[[Workbook], None]) -> bytes:
    workbook = Workbook()
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _assert_search_only_profile(result: ConversionResult) -> None:
    """Every range is a profile label, labelling is total, blocks never straddle."""
    ranges = result.manifest.derivation_ranges
    assert {r.derivation_kind for r in ranges} <= {
        "profile_structure",
        "profile_sample",
    }
    assert ranges[0].start == 0
    assert ranges[-1].end == len(result.document_md)
    for previous, current in zip(ranges, ranges[1:], strict=False):
        assert previous.end == current.start
    assert is_model_free(ranges=ranges)
    blocks = blockize(document_md=result.document_md)
    assert not any(block_eligibility(blocks=blocks, ranges=ranges))
    for block in blocks:
        owners = [
            r for r in ranges if r.start < block.char_end and block.char_start < r.end
        ]
        assert len(owners) == 1, result.document_md[block.char_start : block.char_end]
    coverage = result.manifest.coverage
    assert coverage.policy == "profile"
    assert coverage.complete is False
    assert coverage.gaps[0].startswith("rows are not represented")


def _section(result: ConversionResult, kind: str) -> str:
    return "".join(
        result.document_md[r.start : r.end]
        for r in result.manifest.derivation_ranges
        if r.derivation_kind == kind
    )


# --- routing ------------------------------------------------------------


def test_stock_routes_send_data_files_to_their_profiles() -> None:
    """Data files route by default; ods only where LibreOffice is installed."""
    routes = STOCK_CONVERSION_ROUTE_NAMES
    assert routes[_XLSX] == "spreadsheet"
    assert routes[_XLS] == "spreadsheet"
    assert routes["text/csv"] == "table"
    assert routes["text/tab-separated-values"] == "table"
    for mime in (
        _PARQUET,
        _ARROW,
        _SQLITE,
        "application/x-spss-sav",
        "application/x-spss-por",
        "application/x-sas-xport",
        "application/x-sas-data",
        "application/x-stata-dta",
    ):
        assert routes[mime] == "dataset", mime
    assert _ODS not in stock_route_names(libreoffice_available=False)
    assert stock_route_names(libreoffice_available=True)[_ODS] == "spreadsheet"
    built = build_conversion_routes(route_names=routes)
    assert built[_XLSX].name == "spreadsheet"
    assert built["text/csv"].name == "table"
    assert built[_SQLITE].name == "dataset"


@pytest.mark.parametrize(
    ("name", "mime"),
    [
        ("book.xlsm", _XLSX),
        ("book.xltx", _XLSX),
        ("old.xls", _XLS),
        ("pipes.psv", "text/csv"),
        ("tabs.tab", "text/tab-separated-values"),
        ("frame.feather", _ARROW),
        ("app.db", _SQLITE),
    ],
)
def test_data_extensions_store_their_routed_mime(name: str, mime: str) -> None:
    detected = detect_mime(
        file_name=name, declared_mime="", content=b"x", routed_mimes=()
    )
    assert detected == mime
    assert detected in STOCK_CONVERSION_ROUTE_NAMES


def test_a_spreadsheet_over_its_reading_limit_gets_an_oversized_card(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Over 200 MB a workbook is carded, not profiled (limit shrunk for the test)."""
    assert exceeds_reading_limit(mime=_XLSX, byte_size=200_000_001)
    assert not exceeds_reading_limit(mime=_XLSX, byte_size=200_000_000)
    small = dataclasses.replace(family_for_mime(mime=_XLSX), reading_limit_bytes=10)
    monkeypatch.setattr(
        "rememberstack.core.file_card.family_for_mime", lambda *, mime: small
    )
    result = CardConverter().convert(
        content=b"0" * 11, mime=_XLSX, hints=_hints("huge.xlsx")
    )
    assert "Not read: the file is larger than the 10-byte reading limit" in (
        result.document_md
    )
    assert result.manifest.derivation_ranges[0].derivation_kind == "file_card"


# --- spreadsheets ---------------------------------------------------------


def test_xlsx_profile_lists_sheets_columns_sample_names_and_metadata() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Revenue"
        sheet.append([])
        sheet.append(["Region", "Amount", "Booked", "Closed", "Note"])
        sheet.append(["North", 1200, datetime.datetime(2024, 1, 2), True, "a|b"])
        sheet.append(["South", 3.5, datetime.datetime(2024, 2, 3), False, None])
        for index in range(10):
            sheet.append([f"R{index}", index, None, None, None])
        workbook.create_sheet("Empty")
        workbook.defined_names["TaxRate"] = DefinedName(
            "TaxRate", attr_text="Revenue!$B$3"
        )
        workbook.properties.title = "Q3 sales"
        workbook.properties.creator = "Alice Novak"

    result = SpreadsheetConverter().convert(
        content=_xlsx(build), mime=_XLSX, hints=_hints("q3.xlsx")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert markdown.startswith("# q3.xlsx\n")
    assert "- Family: spreadsheet" in markdown
    assert "- Format: xlsx" in markdown
    assert "- Sheets: 2" in markdown
    assert "## Sheet: Revenue" in markdown
    assert (
        "- Declared dimension: A2:E14 (13 rows × 5 columns, as the file "
        "declares it; not verified)"
    ) in markdown
    assert "- Header row: 2" in markdown
    assert "- Data rows below the header: 12 (from the declared size)" in markdown
    assert "  - Region: text (column A)" in markdown
    assert "  - Amount: number (column B)" in markdown
    assert "  - Booked: date (column C)" in markdown
    assert "  - Closed: boolean (column D)" in markdown
    assert "  - Note: text (column E)" in markdown
    sample = _section(result, "profile_sample")
    assert "### Revenue: first 5 data rows" in sample
    assert "| North | 1200 | 2024-01-02 | True | a\\|b |" in sample
    assert "R2" in sample and "R3" not in markdown
    assert "## Sheet: Empty" in markdown
    assert "Header row: none (the sheet is empty)" in markdown
    names = result.document_md[result.manifest.derivation_ranges[-1].start :]
    assert names.startswith("\n## Defined names (1)")
    assert "- `TaxRate`: Revenue!$B$3" in names
    assert result.manifest.derivation_ranges[-1].evidence_mode == "source_expression"
    assert result.metadata is not None
    assert result.metadata.title == "Q3 sales"
    assert result.metadata.authors[0].name == "Alice Novak"
    assert result.metadata.created_at is not None
    assert result.metadata.created_at.tzinfo is datetime.UTC


def test_xlsx_caps_columns_sample_width_and_cell_length() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Wide"
        sheet.append([f"col{index}" for index in range(120)])
        sheet.append(["x" * 200] + list(range(119)))

    result = SpreadsheetConverter().convert(
        content=_xlsx(build), mime=_XLSX, hints=_hints("wide.xlsx")
    )
    _assert_search_only_profile(result)
    structure = _section(result, "profile_structure")
    assert "- Columns (120):" in structure
    assert "  - col99: integer (column CV)" in structure
    assert "col100:" not in structure
    assert "  - …and 20 more columns" in structure
    sample = _section(result, "profile_sample")
    header = next(line for line in sample.splitlines() if line.startswith("| col0"))
    assert header.count("|") == 31
    assert "The first 30 of 120 columns are shown." in sample
    assert "x" * 79 + "…" in sample
    assert "x" * 80 not in sample


def test_xlsx_describes_fifty_sheets_and_lists_the_rest_by_name() -> None:
    def build(workbook: Workbook) -> None:
        first = workbook.active
        assert first is not None
        first.title = "S00"
        for index in range(1, 55):
            workbook.create_sheet(f"S{index:02d}")
        for index in range(60):
            workbook.defined_names[f"N{index:02d}"] = DefinedName(
                f"N{index:02d}", attr_text="S00!$A$1"
            )

    result = SpreadsheetConverter().convert(
        content=_xlsx(build), mime=_XLSX, hints=_hints("many.xlsx")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "- Sheets: 55" in markdown
    assert "## Sheet: S49" in markdown
    assert "## Sheet: S50" not in markdown
    assert "## Other sheets (5)" in markdown
    assert "Listed by name only: `S50`, `S51`, `S52`, `S53`, `S54`" in markdown
    assert "## Defined names (60)" in markdown
    assert "- `N49`: S00!$A$1" in markdown
    assert "`N50`" not in markdown
    assert "- …and 10 more defined names" in markdown


def test_large_xlsx_is_profiled_from_dimensions_only(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Above the sample limit: sheet list and dimensions, no columns or rows."""

    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["Secret header"])
        sheet.append(["secret value"])

    monkeypatch.setattr(spreadsheet_module, "SAMPLE_LIMIT_BYTES", 10)
    result = SpreadsheetConverter().convert(
        content=_xlsx(build), mime=_XLSX, hints=_hints("big.xlsx")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "## Sheet: Data" in markdown
    assert "- Declared dimension: A1:A2 (2 rows × 1 columns" in markdown
    assert "Sample: not read — the file is larger than 10 bytes" in markdown
    assert "Secret" not in markdown and "secret" not in markdown
    assert "Columns" not in markdown
    assert _section(result, "profile_sample") == ""
    assert result.manifest.coverage.gaps[1].startswith(
        "column lists and sample rows are not read: the file is larger than"
    )


def test_xls_profile_reads_the_legacy_workbook() -> None:
    content = (_FIXTURES / "tiny_budget.xls").read_bytes()
    result = SpreadsheetConverter().convert(
        content=content, mime=_XLS, hints=_hints("budget.xls")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "- Format: xls" in markdown
    assert "- Sheets: 2" in markdown
    assert "- Size: 4 rows × 4 columns" in markdown
    assert "  - Cost: number (column B)" in markdown
    assert "  - Paid: boolean (column C)" in markdown
    assert "  - Due: date (column D)" in markdown
    assert "| Rent | 1200.0 | True | 2024-01-01 |" in markdown
    assert "## Sheet: Notes" in markdown


@pytest.mark.parametrize("mime", [_XLSX, _XLS])
def test_a_corrupt_workbook_fails_with_a_typed_reason(mime: str) -> None:
    with pytest.raises(ConversionError, match="not a readable Excel"):
        SpreadsheetConverter().convert(
            content=b"not a workbook at all", mime=mime, hints=_hints("bad.xlsx")
        )


# --- delimited ------------------------------------------------------------


def test_csv_samples_five_records_and_approximates_the_length() -> None:
    """Quoted newlines stay inside one sampled record; the rest is not parsed."""
    lines = ["id,name,comment"]
    for index in range(1, 8):
        lines.append(f'{index},name{index},"first line\nsecond line {index}"')
    content = ("\n".join(lines) + "\n\n").encode()
    result = TableConverter().convert(
        content=content, mime="text/csv", hints=_hints("notes.csv")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "- Delimiter: comma (detected)" in markdown
    assert "- Encoding: UTF-8" in markdown
    # 1 header line + 7 records of 2 lines each; the blank line counts too
    assert "- Lines in file (approximate, includes header): 16" in markdown
    assert "  - id: integer (column 1)" in markdown
    assert "  - comment: text (column 3)" in markdown
    sample = _section(result, "profile_sample")
    assert "### notes.csv: first 5 data rows" in sample
    assert "| 5 | name5 | first line second line 5 |" in sample
    assert "name6" not in markdown


@pytest.mark.parametrize(
    ("name", "content", "delimiter"),
    [
        ("semi.csv", b"a;b\n1;x\n2;y\n", "semicolon (detected)"),
        ("pipes.psv", b"a|b\n1|x\n2|y\n", "pipe (detected)"),
        ("tabs.tsv", b"a\tb\n1\tx\n2\ty\n", "tab (detected)"),
        ("single.tsv", b"only\n1\n2\n", "tab (from the file type)"),
        ("single.csv", b"only\n1\n2\n", "comma (from the file type)"),
    ],
)
def test_delimited_dialects_are_sniffed_or_taken_from_the_extension(
    name: str, content: bytes, delimiter: str
) -> None:
    result = TableConverter().convert(
        content=content, mime="text/csv", hints=_hints(name)
    )
    assert f"- Delimiter: {delimiter}" in result.document_md
    assert "- Data rows below the header: 2" in result.document_md


def test_csv_encodings_bom_utf16_and_latin1_fallback() -> None:
    bom = TableConverter().convert(
        content="﻿name,city\nZoë,Brno\n".encode(),
        mime="text/csv",
        hints=_hints("bom.csv"),
    )
    assert "  - name: text (column 1)" in bom.document_md
    utf16 = TableConverter().convert(
        content="name,city\nZoë,Brno\n".encode("utf-16"),
        mime="text/csv",
        hints=_hints("wide.csv"),
    )
    assert "- Encoding: UTF-16" in utf16.document_md
    assert "| Zoë | Brno |" in utf16.document_md
    latin = TableConverter().convert(
        content="name,city\nZoë,Brno\n".encode("latin-1"),
        mime="text/csv",
        hints=_hints("latin.csv"),
    )
    _assert_search_only_profile(latin)
    assert "- Encoding: Latin-1" in latin.document_md
    assert "| Zoë | Brno |" in latin.document_md
    assert "the file is not valid UTF-8; it was read as Latin-1" in (
        latin.manifest.coverage.gaps
    )


def test_csv_header_is_the_first_non_empty_record() -> None:
    result = TableConverter().convert(
        content=b"\n,,\nx,y,z\n1,2,3\n", mime="text/csv", hints=_hints("late.csv")
    )
    assert "  - x: integer (column 1)" in result.document_md
    assert "- Data rows below the header: 1" in result.document_md


def test_csv_over_its_time_limit_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(table_module, "TIME_LIMIT_SECONDS", -1.0)
    with pytest.raises(ConversionError, match="time limit"):
        TableConverter().convert(content=b"a,b\n1,2\n", mime="text/csv")


def test_type_inference_from_sample_values() -> None:
    assert infer_type(values=["1", "-2"], parse_text=True) == "integer"
    assert infer_type(values=["1", "2.5"], parse_text=True) == "number"
    assert infer_type(values=["2024-01-31", ""], parse_text=True) == "date"
    assert infer_type(values=["true", "FALSE"], parse_text=True) == "boolean"
    assert infer_type(values=["007", "1"], parse_text=True) == "text"
    assert infer_type(values=["", None], parse_text=True) == "empty"
    assert infer_type(values=["12"], parse_text=False) == "text"
    assert infer_type(values=[1.0, 2], parse_text=False) == "integer"


# --- datasets -------------------------------------------------------------


def _frame_table() -> pa.Table:
    return pa.table(
        {
            "id": list(range(1, 8)),
            "score": [0.5 * index for index in range(7)],
            "label": [f"row{index}" for index in range(7)],
            "day": [datetime.date(2024, 1, index + 1) for index in range(7)],
        }
    )


def test_parquet_profile_reads_footer_counts_and_five_rows() -> None:
    buffer = io.BytesIO()
    pq.write_table(_frame_table(), buffer)
    result = DatasetConverter().convert(
        content=buffer.getvalue(), mime=_PARQUET, hints=_hints("scores.parquet")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "- Family: dataset" in markdown
    assert "- Rows: 7" in markdown
    assert "  - id: integer (column 1; declared int64)" in markdown
    assert "  - label: text (column 3; declared string)" in markdown
    assert "  - day: date (column 4; declared date32[day])" in markdown
    assert "| 5 | 2.0 | row4 | 2024-01-05 |" in markdown
    assert "row5" not in markdown


@pytest.mark.parametrize("form", ["file", "stream", "feather"])
def test_arrow_profile_reads_file_stream_and_feather(form: str) -> None:
    table = _frame_table()
    buffer = io.BytesIO()
    if form == "feather":
        feather.write_feather(table, buffer, chunksize=3)
    else:
        opener = pa.ipc.new_file if form == "file" else pa.ipc.new_stream
        with opener(buffer, table.schema) as writer:
            for batch in table.to_batches(max_chunksize=3):
                writer.write_batch(batch)
    result = DatasetConverter().convert(
        content=buffer.getvalue(), mime=_ARROW, hints=_hints("frame.arrow")
    )
    _assert_search_only_profile(result)
    assert "- Rows: not recorded in the" in result.document_md
    assert "| 5 | 2.0 | row4 | 2024-01-05 |" in result.document_md
    assert "row5" not in result.document_md


@pytest.mark.parametrize(
    ("extension", "mime", "writer", "rows"),
    [
        ("sav", "application/x-spss-sav", pyreadstat.write_sav, "- Rows: 7"),
        ("dta", "application/x-stata-dta", pyreadstat.write_dta, "- Rows: 7"),
        (
            "xpt",
            "application/x-sas-xport",
            pyreadstat.write_xport,
            "- Rows: not recorded in the file header",
        ),
        (
            "por",
            "application/x-spss-por",
            pyreadstat.write_por,
            "- Rows: not recorded in the file header",
        ),
    ],
)
def test_statistical_profiles_read_the_header_and_five_rows(
    tmp_path: Path, extension: str, mime: str, writer: Callable[..., None], rows: str
) -> None:
    frame = pd.DataFrame(
        {
            "ID": list(range(1, 8)),
            "LABEL": [f"row{index}" for index in range(7)],
            "DAY": [datetime.date(2024, 1, index + 1) for index in range(7)],
        }
    )
    path = tmp_path / f"survey.{extension}"
    writer(frame, str(path))
    result = DatasetConverter().convert(
        content=path.read_bytes(), mime=mime, hints=_hints(path.name)
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert rows in markdown
    assert "  - ID: integer (column 1; declared" in markdown
    assert "  - LABEL: text (column 2; declared" in markdown
    assert "  - DAY: date (column 3; declared" in markdown
    assert "row4" in markdown and "row5" not in markdown


def test_sqlite_profile_lists_tables_types_counts_and_samples(tmp_path: Path) -> None:
    path = tmp_path / "app.db"
    with sqlite3.connect(path) as connection:
        connection.execute("CREATE TABLE users (id INTEGER, name TEXT, score REAL)")
        connection.executemany(
            "INSERT INTO users VALUES (?, ?, ?)",
            [(index, f"user{index}", index / 2) for index in range(9)],
        )
        connection.execute('CREATE TABLE "odd ""name""" (blob BLOB)')
        connection.execute('INSERT INTO "odd ""name""" VALUES (x\'0102\')')
    connection.close()
    result = DatasetConverter().convert(
        content=path.read_bytes(), mime=_SQLITE, hints=_hints("app.db")
    )
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "- Tables: 2" in markdown
    assert "## Table: users" in markdown
    assert "- Rows: 9" in markdown
    assert "  - name: text (column 2; declared TEXT)" in markdown
    assert "  - score: number (column 3; declared REAL)" in markdown
    assert "| 4 | user4 | 2.0 |" in markdown
    assert "user5" not in markdown
    assert '## Table: odd "name"' in markdown
    assert "| (2 bytes) |" in markdown


def test_sqlite_describes_fifty_tables(tmp_path: Path) -> None:
    path = tmp_path / "many.sqlite"
    with sqlite3.connect(path) as connection:
        for index in range(53):
            connection.execute(f"CREATE TABLE t{index:02d} (v INTEGER)")
    connection.close()
    result = DatasetConverter().convert(
        content=path.read_bytes(), mime=_SQLITE, hints=_hints("many.sqlite")
    )
    assert "- Tables: 53" in result.document_md
    assert "## Table: t49" in result.document_md
    assert "## Other tables (3)" in result.document_md
    assert "`t50`, `t51`, `t52`" in result.document_md


def test_sqlite_over_its_time_limit_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "slow.sqlite"
    with sqlite3.connect(path) as connection:
        connection.execute("CREATE TABLE t (v INTEGER)")
    connection.close()
    monkeypatch.setattr(dataset_module, "TIME_LIMIT_SECONDS", -1.0)
    with pytest.raises(ConversionError, match="time limit"):
        DatasetConverter().convert(content=path.read_bytes(), mime=_SQLITE)


@pytest.mark.parametrize(
    ("mime", "reason"),
    [
        (_PARQUET, "not a readable Parquet file"),
        (_ARROW, "not a readable Arrow file"),
        (_SQLITE, "not a readable SQLite file"),
        ("application/x-spss-sav", "not a readable statistical data file"),
    ],
)
def test_a_corrupt_dataset_fails_with_a_typed_reason(mime: str, reason: str) -> None:
    with pytest.raises(ConversionError, match=reason):
        DatasetConverter().convert(
            content=b"this is not the declared format" * 10, mime=mime
        )


# --- review follow-ups ----------------------------------------------------


def _rewrite_sheet_xml(content: bytes, rewrite: Callable[[str], str]) -> bytes:
    """Copy an .xlsx, passing the first worksheet's XML through ``rewrite``."""
    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = rewrite(data.decode()).encode()
            target.writestr(info, data)
    return buffer.getvalue()


def _eight_rows(workbook: Workbook) -> None:
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["name", "value"])
    for index in range(7):
        sheet.append([f"row{index}", index])


def test_an_understated_dimension_does_not_cut_the_sample() -> None:
    content = _rewrite_sheet_xml(
        _xlsx(_eight_rows),
        lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:A2"', xml),
    )
    result = SpreadsheetConverter().convert(content=content, mime=_XLSX)
    markdown = result.document_md
    assert "- Declared dimension: A1:A2 (2 rows × 1 columns" in markdown
    assert "  - value: integer (column B)" in markdown
    assert "| row4 | 4 |" in markdown
    assert "- Data rows below the header: at least 5" in markdown


def test_an_unsized_sheet_is_still_sampled() -> None:
    content = _rewrite_sheet_xml(
        _xlsx(_eight_rows), lambda xml: re.sub(r"<dimension [^>]*/>", "", xml)
    )
    result = SpreadsheetConverter().convert(content=content, mime=_XLSX)
    markdown = result.document_md
    assert "- Declared dimension: not recorded" in markdown
    assert "| row4 | 4 |" in markdown
    assert "- Data rows below the header: at least 5" in markdown


def test_an_overstated_dimension_reports_the_rows_actually_read() -> None:
    content = _rewrite_sheet_xml(
        _xlsx(_eight_rows),
        lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:B3"', xml),
    )
    small = _rewrite_sheet_xml(
        content, lambda xml: re.sub(r'<row r="([4-8])".*?</row>', "", xml, flags=re.S)
    )
    result = SpreadsheetConverter().convert(content=small, mime=_XLSX)
    assert "- Data rows below the header: 2" in result.document_md


def test_formula_cells_without_cached_values_show_their_formula() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["a", "b", "total"])
        sheet.append([1, 2, "=A2+B2"])
        sheet.append([3, 4, "=A3+B3"])

    result = SpreadsheetConverter().convert(content=_xlsx(build), mime=_XLSX)
    markdown = result.document_md
    assert "| 1 | 2 | =A2+B2 |" in markdown
    assert "  - total: empty (column C)" in markdown


def test_hidden_sheets_are_marked() -> None:
    def build(workbook: Workbook) -> None:
        first = workbook.active
        assert first is not None
        first.title = "Shown"
        workbook.create_sheet("Secret").sheet_state = "hidden"
        workbook.create_sheet("Deep").sheet_state = "veryHidden"

    markdown = (
        SpreadsheetConverter().convert(content=_xlsx(build), mime=_XLSX).document_md
    )
    shown = markdown[
        markdown.index("## Sheet: Shown") : markdown.index("## Sheet: Secret")
    ]
    assert "Visibility" not in shown
    assert "## Sheet: Secret\n\n- Visibility: hidden" in markdown
    assert "## Sheet: Deep\n\n- Visibility: very hidden" in markdown


def test_a_large_shared_strings_table_skips_the_sample(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(spreadsheet_module, "SHARED_STRINGS_LIMIT_BYTES", 10)
    buffer = io.BytesIO(_xlsx(_eight_rows))
    with zipfile.ZipFile(buffer, "a") as archive:
        # openpyxl writes inline strings; add the table another writer would
        archive.writestr("xl/sharedStrings.xml", "<sst>" + "x" * 100 + "</sst>")
    result = SpreadsheetConverter().convert(content=buffer.getvalue(), mime=_XLSX)
    markdown = result.document_md
    assert "Sample: not read — its shared-strings table expands to" in markdown
    assert "row0" not in markdown
    assert "- Declared dimension: A1:B8" in markdown


def test_a_large_xls_lists_sheet_names_without_loading_sheets(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(spreadsheet_module, "XLS_SHEET_LOAD_LIMIT_BYTES", 10)

    def refuse(*args: object, **kwargs: object) -> None:
        raise AssertionError("a sheet was loaded")

    monkeypatch.setattr("xlrd.book.Book.sheet_by_index", refuse)
    content = (_FIXTURES / "tiny_budget.xls").read_bytes()
    result = SpreadsheetConverter().convert(content=content, mime=_XLS)
    _assert_search_only_profile(result)
    markdown = result.document_md
    assert "## Sheet: Budget" in markdown and "## Sheet: Notes" in markdown
    assert "Size: not read (the sheet was not loaded)" in markdown
    assert "Rent" not in markdown


def test_spreadsheet_over_its_time_limit_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(spreadsheet_module, "TIME_LIMIT_SECONDS", -1.0)
    with pytest.raises(ConversionError, match="time limit"):
        SpreadsheetConverter().convert(content=_xlsx(_eight_rows), mime=_XLSX)


def test_csv_field_over_the_size_limit_fails_clearly(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(table_module, "FIELD_SIZE_LIMIT", 100)
    content = b"a,b\n1," + b"x" * 500 + b"\n"
    with pytest.raises(ConversionError, match="longer than the 100-character limit"):
        TableConverter().convert(content=content, mime="text/csv")
    monkeypatch.undo()
    result = TableConverter().convert(content=content, mime="text/csv")
    assert "x" * 79 + "…" in result.document_md


def test_arrow_stops_reading_once_the_sample_is_full() -> None:
    """Batches after the fifth row are never decoded (a truncated tail is fine)."""
    table = _frame_table()
    buffer = io.BytesIO()
    with pa.ipc.new_stream(buffer, table.schema) as writer:
        for batch in table.to_batches(max_chunksize=3):
            writer.write_batch(batch)
        writer.write_batch(table.to_batches()[0])
    content = buffer.getvalue()
    result = DatasetConverter().convert(content=content[:-200], mime=_ARROW)
    assert "| 5 | 2.0 | row4 | 2024-01-05 |" in result.document_md


def test_sqlite_generated_columns_views_and_bounded_values(tmp_path: Path) -> None:
    path = tmp_path / "gen.sqlite"
    with sqlite3.connect(path) as connection:
        connection.execute(
            "CREATE TABLE t (a INTEGER, body TEXT, "
            "twice INTEGER GENERATED ALWAYS AS (a * 2) VIRTUAL)"
        )
        connection.execute("INSERT INTO t (a, body) VALUES (3, ?)", ("y" * 5000,))
        connection.execute("CREATE VIEW doubled AS SELECT twice FROM t")
    connection.close()
    markdown = (
        DatasetConverter().convert(content=path.read_bytes(), mime=_SQLITE).document_md
    )
    assert "  - twice: integer (column 3; declared INTEGER)" in markdown
    assert "| 3 | " + "y" * 79 + "… | 6 |" in markdown
    assert "- Views: 1 (`doubled`)" in markdown


def test_a_views_only_sqlite_database_names_its_views(tmp_path: Path) -> None:
    path = tmp_path / "views.sqlite"
    with sqlite3.connect(path) as connection:
        connection.execute("CREATE VIEW answer AS SELECT 42 AS value")
    connection.close()
    markdown = (
        DatasetConverter().convert(content=path.read_bytes(), mime=_SQLITE).document_md
    )
    assert "- Tables: 0" in markdown
    assert "- Views: 1 (`answer`)" in markdown


def test_nan_is_ignored_when_inferring_types() -> None:
    assert infer_type(values=[float("nan"), None], parse_text=False) == "empty"
    assert infer_type(values=[float("nan"), 1.5], parse_text=False) == "number"


def test_ods_is_converted_by_libreoffice_then_profiled(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The .ods bytes go to LibreOffice; the xlsx it returns is profiled."""
    calls: list[tuple[bytes, str, str]] = []

    def fake(*, content: bytes, source_extension: str, target: str) -> bytes:
        calls.append((content, source_extension, target))
        return _xlsx(_eight_rows)

    monkeypatch.setattr(spreadsheet_module, "convert_with_libreoffice", fake)
    result = SpreadsheetConverter().convert(
        content=b"ods bytes", mime=_ODS, hints=_hints("sheet.ods")
    )
    _assert_search_only_profile(result)
    assert calls == [(b"ods bytes", "ods", "xlsx")]
    assert "- Format: ods" in result.document_md
    assert "- Byte size: 9 bytes" in result.document_md
    assert "| row4 | 4 |" in result.document_md
    assert [c.name for c in result.manifest.components] == [
        "spreadsheet",
        "libreoffice",
    ]


@pytest.mark.skipif(
    shutil.which("soffice") is None, reason="LibreOffice (soffice) is not installed"
)
def test_ods_profile_with_real_libreoffice(tmp_path: Path) -> None:
    from rememberstack.adapters.converters.libreoffice import convert_with_libreoffice

    ods = convert_with_libreoffice(
        content=_xlsx(_eight_rows),
        source_extension="xlsx",
        target="ods",  # type: ignore[arg-type]
    )
    result = SpreadsheetConverter().convert(
        content=ods, mime=_ODS, hints=_hints("sheet.ods")
    )
    assert "| row4 | 4 |" in result.document_md


def test_csv_parsing_stops_after_the_sample() -> None:
    """A malformed record after the sample is never parsed."""
    content = b"a,b\n" + b"1,2\n" * 5 + b'3,"unterminated\n' + b"x" * 50
    result = TableConverter().convert(content=content, mime="text/csv")
    assert "| 1 | 2 |" in result.document_md
    assert "- Lines in file (approximate, includes header): 8" in result.document_md


def test_a_formula_only_header_and_rows_are_not_empty() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        assert sheet is not None
        sheet.append(['=CONCAT("q", 1)', '=CONCAT("q", 2)'])
        sheet.append(["=1+1", "=2+2"])
        sheet.append(["=3+3", "=4+4"])

    markdown = (
        SpreadsheetConverter().convert(content=_xlsx(build), mime=_XLSX).document_md
    )
    assert "- Header row: 1" in markdown
    assert '  - =CONCAT("q", 1): empty (column A)' in markdown
    assert "| =1+1 | =2+2 |" in markdown
    assert "| =3+3 | =4+4 |" in markdown
    assert "- Data rows below the header: 2" in markdown
