"""The D138 document-family routes: office, PDF, HTML, e-book, email, notebook.

Every fixture is generated here, small and deterministic: python-pptx decks,
hand-assembled OOXML and PDF bytes, stdlib email messages and notebook JSON.
LibreOffice conversions run only where ``soffice`` is installed.
"""

from datetime import datetime
from datetime import timezone
from email.message import EmailMessage
import io
import json
from pathlib import Path
import shutil
import subprocess
import time
import zipfile

from pptx import Presentation
from pptx.util import Inches
import pytest

from rememberstack.adapters.converters import build_conversion_routes
from rememberstack.adapters.converters import libreoffice
from rememberstack.adapters.converters import pdf
from rememberstack.adapters.converters import time_limit
from rememberstack.adapters.converters import zip_budget
from rememberstack.adapters.converters.card import CardConverter
from rememberstack.adapters.converters.email_message import EmailConverter
from rememberstack.adapters.converters.markitdown import MarkitdownConverter
from rememberstack.adapters.converters.notebook import NotebookConverter
from rememberstack.adapters.converters.office import OfficeConverter
from rememberstack.adapters.converters.pdf import PdfConverter
from rememberstack.core import ConversionRouter
from rememberstack.core.extraction_eligibility import INELIGIBLE_DERIVATION_KINDS
from rememberstack.core.format_registry import exceeds_reading_limit
from rememberstack.core.format_registry import stock_route_names
from rememberstack.model import ConversionCoverage
from rememberstack.model import ConversionError
from rememberstack.model import ConversionResult
from rememberstack.model import ConverterManifest
from rememberstack.model import FileHints
from rememberstack.model import ManifestComponent
from rememberstack.model import PageLocator
from rememberstack.model import SourceMapEntry
from rememberstack.model.document_metadata import DocumentMetadata
from rememberstack.model.document_metadata import DocumentPerson

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_SOFFICE_MISSING = "LibreOffice (soffice) is not installed on this machine"


# --- office: presentations ---------------------------------------------------


def test_presentation_slides_get_their_own_page_locators() -> None:
    """Each slide is one ``## Slide N`` section mapped to ``page`` = N."""
    result = OfficeConverter().convert(content=_pptx(), mime=_PPTX)
    markdown = result.document_md
    assert "## Slide 1: Roadmap review" in markdown
    assert "Launch in October." in markdown
    assert "Speaker notes:\n\nMention the budget." in markdown
    assert "## Slide 2" in markdown
    assert "| Region | Revenue |" in markdown
    assert "| North | 1200 |" in markdown
    assert result.source_map is not None
    pages = [entry.locators[0] for entry in result.source_map]
    assert pages == [
        PageLocator(page=1, precision="page"),
        PageLocator(page=2, precision="page"),
    ]
    first, second = result.source_map
    assert markdown[first.start : first.end].startswith("## Slide 1")
    assert markdown[second.start : second.end].startswith("## Slide 2")
    _assert_prose(result=result)


def test_presentation_core_properties_become_d134_metadata() -> None:
    """Title, creator, created and modified come from docProps/core.xml."""
    result = OfficeConverter().convert(content=_pptx(), mime=_PPTX)
    assert result.metadata is not None
    assert result.metadata.title == "Q4 roadmap"
    assert result.metadata.authors == (DocumentPerson(name="Alice Novak"),)
    assert result.metadata.created_at == datetime(
        2026, 1, 2, 3, 4, 5, tzinfo=timezone.utc
    )
    assert result.metadata.modified_at == datetime(
        2026, 2, 3, 4, 5, 6, tzinfo=timezone.utc
    )


@pytest.mark.parametrize("kind", ["slideshow", "template"])
def test_slide_shows_and_templates_read_like_a_pptx(kind: str) -> None:
    """ppsx and potx differ only in the main content type; both are read."""
    content = _with_replaced_member(
        content=_pptx(),
        member="[Content_Types].xml",
        old=b"presentationml.presentation.main+xml",
        new=f"presentationml.{kind}.main+xml".encode(),
    )
    result = OfficeConverter().convert(content=content, mime=_PPTX)
    assert "## Slide 1: Roadmap review" in result.document_md


def test_corrupt_presentation_fails_with_a_typed_error() -> None:
    """Bytes that are not an OOXML package fail; nothing falls back."""
    with pytest.raises(ConversionError, match="not a readable zip package"):
        OfficeConverter().convert(content=b"not a zip at all", mime=_PPTX)
    broken = _with_replaced_member(
        content=_pptx(),
        member="ppt/presentation.xml",
        old=b"<p:presentation",
        new=b"<p:broken",
    )
    with pytest.raises(ConversionError, match="presentation could not be read"):
        OfficeConverter().convert(content=broken, mime=_PPTX)


def test_slide_title_whitespace_is_collapsed_in_the_heading() -> None:
    """A title with line breaks still makes a one-line ``## Slide`` heading."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Budget\nreview   2026"
    buffer = io.BytesIO()
    deck.save(buffer)
    result = OfficeConverter().convert(content=buffer.getvalue(), mime=_PPTX)
    assert "## Slide 1: Budget review 2026\n" in result.document_md


def test_zip_packages_over_the_expansion_budget_fail(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Declared uncompressed sizes are checked before OOXML/EPUB parsing."""
    monkeypatch.setattr(zip_budget, "MAX_UNCOMPRESSED_MEMBER_BYTES", 100)
    with pytest.raises(ConversionError, match="member larger than"):
        OfficeConverter().convert(content=_pptx(), mime=_PPTX)
    monkeypatch.setattr(zip_budget, "MAX_UNCOMPRESSED_MEMBER_BYTES", 10**9)
    monkeypatch.setattr(zip_budget, "MAX_UNCOMPRESSED_TOTAL_BYTES", 100)
    with pytest.raises(ConversionError, match="expands to more than"):
        OfficeConverter().convert(content=_docx(), mime=_DOCX)
    with pytest.raises(ConversionError, match="expands to more than"):
        MarkitdownConverter().convert(content=_epub(), mime="application/epub+zip")


def test_damaged_zip_member_is_a_typed_failure() -> None:
    """A member whose compressed bytes are corrupt fails as ConversionError."""
    content = bytearray(_pptx())
    start = content.index(b"ppt/presentation.xml") + len(b"ppt/presentation.xml")
    content[start + 20 : start + 60] = b"\xff" * 40
    with pytest.raises(ConversionError):
        OfficeConverter().convert(content=bytes(content), mime=_PPTX)


# --- office: Word ------------------------------------------------------------


def test_word_document_is_prose_with_core_metadata() -> None:
    """A docx is rendered by markitdown; its core properties fill D134."""
    result = OfficeConverter().convert(content=_docx(), mime=_DOCX)
    assert "Ship the converter." in result.document_md
    _assert_prose(result=result)
    assert result.metadata is not None
    assert result.metadata.title == "Plan"
    assert result.metadata.authors == (DocumentPerson(name="Bob Smith"),)
    assert result.metadata.created_at == datetime(
        2025, 5, 6, 7, 8, 9, tzinfo=timezone.utc
    )


def test_word_core_properties_with_a_dtd_are_skipped_not_expanded() -> None:
    """defusedxml refuses entity declarations; only the metadata is lost."""
    evil = (
        b'<?xml version="1.0"?><!DOCTYPE t [<!ENTITY a "boom">]>'
        b'<cp:coreProperties xmlns:cp="x" xmlns:dc="http://purl.org/dc/elements/1.1/">'
        b"<dc:title>&a;</dc:title></cp:coreProperties>"
    )
    result = OfficeConverter().convert(content=_docx(core_xml=evil), mime=_DOCX)
    assert "Ship the converter." in result.document_md
    assert result.metadata is None
    assert result.warnings


def test_corrupt_word_document_fails_with_a_typed_error() -> None:
    """A zip without a Word body is a typed failure, not an empty reading."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        package.writestr("hello.txt", "no word body here")
    with pytest.raises(ConversionError):
        OfficeConverter().convert(content=buffer.getvalue(), mime=_DOCX)


# --- office: LibreOffice -----------------------------------------------------


def test_legacy_office_formats_route_only_with_libreoffice() -> None:
    """Without soffice, doc/odt/rtf/ppt/odp have no route and park (D117)."""
    legacy = (
        "application/msword",
        "application/vnd.oasis.opendocument.text",
        "application/rtf",
        "application/vnd.ms-powerpoint",
        "application/vnd.oasis.opendocument.presentation",
    )
    without = stock_route_names(libreoffice_available=False)
    with_it = stock_route_names(libreoffice_available=True)
    for mime in legacy:
        assert mime not in without, mime
        assert with_it[mime] == "office", mime
    assert without[_DOCX] == with_it[_DOCX] == "office"
    ods = "application/vnd.oasis.opendocument.spreadsheet"
    assert ods not in without
    assert with_it[ods] == "spreadsheet"


def test_legacy_format_without_soffice_is_a_typed_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A deployment routing rtf explicitly without LibreOffice fails, typed."""
    monkeypatch.setattr(libreoffice.shutil, "which", lambda _name: None)
    with pytest.raises(ConversionError, match="not installed"):
        OfficeConverter().convert(content=b"{\\rtf1 hi}", mime="application/rtf")


def test_libreoffice_is_killed_at_the_time_limit(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """A hanging soffice is killed and the conversion fails, typed."""
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\nsleep 30\n")
    fake.chmod(0o755)
    monkeypatch.setattr(libreoffice.shutil, "which", lambda _name: str(fake))
    monkeypatch.setattr(libreoffice, "CONVERTER_TIME_LIMIT_S", 0.5)
    started = time.monotonic()
    with pytest.raises(ConversionError, match="did not convert"):
        libreoffice.convert_with_libreoffice(
            content=b"{\\rtf1 hi}", source_extension="rtf", target="docx"
        )
    assert time.monotonic() - started < 10


@pytest.mark.skipif(shutil.which("soffice") is None, reason=_SOFFICE_MISSING)
def test_libreoffice_converts_rtf_to_a_word_reading() -> None:
    """An RTF file goes through soffice → docx → markitdown."""
    rtf = rb"{\rtf1\ansi {\info{\title Memo}{\author Carol}}Budget approved.\par}"
    result = OfficeConverter().convert(content=rtf, mime="application/rtf")
    assert "Budget approved." in result.document_md
    assert "libreoffice" in [c.name for c in result.manifest.components]
    _assert_prose(result=result)


@pytest.mark.skipif(shutil.which("soffice") is None, reason=_SOFFICE_MISSING)
def test_libreoffice_converts_odp_to_slides(tmp_path: Path) -> None:
    """An odp (made from a pptx by soffice) is read slide by slide."""
    source = tmp_path / "deck.pptx"
    source.write_bytes(_pptx())
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "odp", "--outdir", str(tmp_path)]
        + [str(source)],
        check=True,
        capture_output=True,
        timeout=120,
    )
    content = (tmp_path / "deck.odp").read_bytes()
    result = OfficeConverter().convert(
        content=content, mime="application/vnd.oasis.opendocument.presentation"
    )
    assert "Roadmap review" in result.document_md
    assert result.source_map is not None
    assert result.source_map[0].locators[0] == PageLocator(page=1, precision="page")


@pytest.mark.skipif(shutil.which("soffice") is None, reason=_SOFFICE_MISSING)
def test_libreoffice_rejects_corrupt_input() -> None:
    """Garbage declared as .doc produces no output: a typed failure."""
    with pytest.raises(ConversionError):
        OfficeConverter().convert(
            content=b"\x00\x01garbage" * 10, mime="application/msword"
        )


# --- PDF ---------------------------------------------------------------------


def test_every_pdf_page_uses_ocr_even_when_text_is_selectable() -> None:
    """Born-digital, scanned and mixed pages all take the same OCR path."""
    ocr = _FakeOcr(pages=(1, 2, 3))
    content = _pdf(pages=["Selectable", None, "More selectable"])
    result = PdfConverter(ocr=ocr).convert(content=content, mime="application/pdf")
    assert ocr.calls == 1
    assert result.document_md == "OCR text\n"
    assert result.source_map is not None
    assert {entry.locators[0].page for entry in result.source_map} == {1, 2, 3}  # pyright: ignore[reportAttributeAccessIssue]
    assert result.manifest.coverage.complete


def test_pdf_missing_ocr_page_fails_without_partial_reading() -> None:
    """A provider response gap is a typed failure, not a partial success."""
    with pytest.raises(ConversionError, match="PDF OCR incomplete"):
        PdfConverter(ocr=_FakeOcr(pages=(1, 3))).convert(
            content=_pdf(pages=["One", None, "Three"]), mime="application/pdf"
        )


def test_pdf_info_dictionary_becomes_d134_metadata() -> None:
    """Structural Info metadata survives the OCR reading."""
    content = _pdf(
        pages=["Body"],
        info={
            "Title": "Audit 2025",
            "Author": "Dana Lee",
            "CreationDate": "D:20250301120000+01'00'",
            "ModDate": "D:20250302",
        },
    )
    result = PdfConverter(ocr=_FakeOcr()).convert(
        content=content, mime="application/pdf"
    )
    assert result.metadata is not None
    assert result.metadata.title == "Audit 2025"
    assert result.metadata.authors == (DocumentPerson(name="Dana Lee"),)
    assert result.metadata.created_at == datetime(2025, 3, 1, 11, tzinfo=timezone.utc)
    assert result.metadata.modified_at == datetime(2025, 3, 2, tzinfo=timezone.utc)


def test_ocr_metadata_is_kept_when_the_pdf_declares_none() -> None:
    """Without an Info dictionary the OCR route's own metadata stays."""
    ocr = _FakeOcr(metadata=DocumentMetadata(title="From OCR"))
    result = PdfConverter(ocr=ocr).convert(
        content=_pdf(pages=[None]), mime="application/pdf"
    )
    assert result.metadata == DocumentMetadata(title="From OCR")


def test_pdf_fails_when_pdfium_is_still_held(monkeypatch: pytest.MonkeyPatch) -> None:
    """A stuck earlier page-count operation fails the next one, typed."""
    monkeypatch.setattr(time_limit, "CONVERTER_TIME_LIMIT_S", 0.05)
    assert pdf._PDFIUM_LOCK.acquire(timeout=1)  # pyright: ignore[reportPrivateUsage]
    try:
        with pytest.raises(ConversionError, match="still held"):
            PdfConverter(ocr=_FakeOcr()).convert(
                content=_pdf(pages=["x"]), mime="application/pdf"
            )
    finally:
        pdf._PDFIUM_LOCK.release()  # pyright: ignore[reportPrivateUsage]


def test_pdf_route_parks_without_key_and_rejects_non_ocr_overlay(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The stock PDF entry only composes when the OCR provider is ready."""
    monkeypatch.delenv("REMEMBERSTACK_MISTRAL_OCR_API_KEY", raising=False)
    assert "application/pdf" not in build_conversion_routes(
        route_names={"application/pdf": "pdf"}
    )
    monkeypatch.setenv("REMEMBERSTACK_MISTRAL_OCR_API_KEY", "test-key")
    routes = build_conversion_routes(route_names={"application/pdf": "pdf"})
    assert routes["application/pdf"].version.startswith("pdf-ocr-2026.09+mistral-ocr")
    with pytest.raises(Exception, match="every-page OCR"):
        build_conversion_routes(route_names={"application/pdf": "card"})


def test_corrupt_pdf_fails_with_a_typed_error() -> None:
    """No PDF body, or a header over garbage, fails before OCR."""
    with pytest.raises(ConversionError, match="not a PDF"):
        PdfConverter(ocr=_FakeOcr()).convert(content=b"hello", mime="application/pdf")
    with pytest.raises(ConversionError, match="not a PDF"):
        PdfConverter(ocr=_FakeOcr()).convert(
            content=b"%PDF-1.4\ngarbage without objects", mime="application/pdf"
        )


def test_oversized_non_pdf_documents_get_a_card() -> None:
    """Office documents above the family limit still get a card."""
    for mime in ("application/pdf", _DOCX, _PPTX, "application/msword"):
        assert exceeds_reading_limit(mime=mime, byte_size=100_000_001), mime
        assert not exceeds_reading_limit(mime=mime, byte_size=100_000_000), mime
    card = CardConverter().convert(
        content=b"\0" * 100_000_001,
        mime=_DOCX,
        hints=FileHints(file_name="huge.docx", source_path="docs/huge.docx"),
    )
    assert "100,000,000-byte reading limit for word files" in card.document_md


# --- HTML and e-book (markitdown) --------------------------------------------


def test_html_and_epub_route_to_markitdown_and_read_as_prose() -> None:
    """The stock table sends HTML and EPUB to markitdown; both are prose."""
    routes = stock_route_names(libreoffice_available=False)
    assert routes["text/html"] == routes["application/epub+zip"] == "markitdown"
    converter = MarkitdownConverter()
    html = converter.convert(
        content=b"<html><head><title>Notes</title></head><body><h1>Hi</h1>"
        b"<p>Launch approved.</p></body></html>",
        mime="text/html",
    )
    assert "Launch approved." in html.document_md
    ebook = converter.convert(content=_epub(), mime="application/epub+zip")
    assert "Chapter one begins." in ebook.document_md
    for result in (html, ebook):
        assert all(
            labeled.derivation_kind not in INELIGIBLE_DERIVATION_KINDS
            for labeled in result.manifest.derivation_ranges
        )


def test_corrupt_epub_fails_with_a_typed_error() -> None:
    """An EPUB that is not a zip is a typed failure."""
    with pytest.raises(ConversionError):
        MarkitdownConverter().convert(content=b"not a zip", mime="application/epub+zip")


# --- email -------------------------------------------------------------------


def test_email_headers_body_and_attachments() -> None:
    """Headers, the plain body and the attachment list; nothing converted."""
    result = EmailConverter().convert(content=_eml(), mime="message/rfc822")
    markdown = result.document_md
    assert markdown.startswith("# Budget approval\n")
    assert "- From: Alice Novak <alice@acme.com>" in markdown
    assert "- To: Bob <bob@acme.com>" in markdown
    assert "- Cc: carol@acme.com" in markdown
    assert "The budget is approved." in markdown
    assert "## Attachments\n\n- budget.xlsx (10 bytes)" in markdown
    assert result.manifest.coverage.gaps == ("attachments are listed, not converted",)
    _assert_prose(result=result)


def test_email_metadata_maps_onto_d134() -> None:
    """Subject, From, To+Cc, Date and the References root."""
    result = EmailConverter().convert(content=_eml(), mime="message/rfc822")
    metadata = result.metadata
    assert metadata is not None
    assert metadata.title == "Budget approval"
    assert metadata.authors == (
        DocumentPerson(name="Alice Novak", address="alice@acme.com"),
    )
    assert metadata.recipients == (
        DocumentPerson(name="Bob", address="bob@acme.com"),
        DocumentPerson(address="carol@acme.com"),
    )
    assert metadata.created_at == datetime(2026, 3, 4, 9, 30, tzinfo=timezone.utc)
    assert metadata.thread_ref == "<root@acme.com>"


def test_html_only_email_is_converted_with_markitdown() -> None:
    """Without a plain part the HTML body is rendered; Message-ID is the thread."""
    message = EmailMessage()
    message["From"] = "alice@acme.com"
    message["Subject"] = "Launch"
    message["Message-ID"] = "<only@acme.com>"
    message.set_content("<p>We launch on <b>Monday</b>.</p>", subtype="html")
    result = EmailConverter().convert(content=message.as_bytes(), mime="message/rfc822")
    assert "We launch on **Monday**." in result.document_md
    assert result.manifest.coverage.complete is True
    assert result.metadata is not None
    assert result.metadata.thread_ref == "<only@acme.com>"


def test_single_part_attachment_message_lists_its_root_part() -> None:
    """A message whose only part is an attachment lists it; not complete."""
    message = EmailMessage()
    message["From"] = "alice@acme.com"
    message["Subject"] = "Report"
    message.set_content(b"%PDF-1.4 data", maintype="application", subtype="pdf")
    message.add_header("Content-Disposition", "attachment", filename="report.pdf")
    result = EmailConverter().convert(content=message.as_bytes(), mime="message/rfc822")
    assert "## Attachments\n\n- report.pdf (13 bytes)" in result.document_md
    assert result.manifest.coverage.complete is False


def test_bytes_without_email_headers_fail() -> None:
    """A file with no From, To, Subject or Date is not an email."""
    with pytest.raises(ConversionError, match="not an email"):
        EmailConverter().convert(content=b"\x00\x01random", mime="message/rfc822")


# --- notebook ----------------------------------------------------------------


def test_notebook_cells_in_order_with_code_ineligible() -> None:
    """Markdown cells are prose, code cells ``code``; outputs are dropped."""
    result = NotebookConverter().convert(
        content=_ipynb(), mime="application/x-ipynb+json"
    )
    markdown = result.document_md
    assert markdown.index("# Churn study") < markdown.index("```python\nimport pandas")
    assert markdown.index("import pandas") < markdown.index("Churn fell by 3%.")
    assert "SECRET OUTPUT" not in markdown
    kinds = [
        (labeled.derivation_kind, markdown[labeled.start : labeled.end].strip())
        for labeled in result.manifest.derivation_ranges
    ]
    assert [kind for kind, _ in kinds] == ["prose", "code", "prose"]
    assert "code" in INELIGIBLE_DERIVATION_KINDS
    assert "prose" not in INELIGIBLE_DERIVATION_KINDS
    assert result.metadata is not None
    assert result.metadata.title == "Churn study"
    assert result.manifest.coverage.gaps == ("outputs of 1 cells dropped",)


def test_notebook_labels_fall_on_block_boundaries() -> None:
    """E0 accepts the envelope: eligibility changes only between blocks."""
    from rememberstack.core import blockize
    from rememberstack.core.extraction_eligibility import block_eligibility

    result = NotebookConverter().convert(
        content=_ipynb(), mime="application/x-ipynb+json"
    )
    ranges = result.manifest.derivation_ranges
    eligibility = block_eligibility(
        blocks=blockize(document_md=result.document_md), ranges=ranges
    )
    assert eligibility == (True, True, False, True)


def test_raw_cells_null_sources_and_other_languages() -> None:
    """Raw cells are other_text; a null source is skipped; c++ fences."""
    notebook = {
        "metadata": {"language_info": {"name": "c++"}},
        "cells": [
            {"cell_type": "markdown", "source": None},
            {"cell_type": "raw", "source": "raw text"},
            {"cell_type": "code", "source": "int main() {}"},
        ],
    }
    result = NotebookConverter().convert(
        content=json.dumps(notebook).encode(), mime="application/x-ipynb+json"
    )
    assert "None" not in result.document_md
    assert "```c++\nint main() {}" in result.document_md
    kinds = [labeled.derivation_kind for labeled in result.manifest.derivation_ranges]
    assert kinds == ["other_text", "code"]


def test_invalid_notebook_fails_with_a_typed_error() -> None:
    """Invalid JSON, or JSON without a cell list, is a typed failure."""
    converter = NotebookConverter()
    with pytest.raises(ConversionError, match="not valid JSON"):
        converter.convert(content=b"{not json", mime="application/x-ipynb+json")
    with pytest.raises(ConversionError, match="no cell list"):
        converter.convert(content=b'{"a": 1}', mime="application/x-ipynb+json")


# --- routing and limits ------------------------------------------------------


def test_every_document_family_routes_through_the_stock_table(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The registry's converter names all build and route by default."""
    monkeypatch.setenv("REMEMBERSTACK_MISTRAL_OCR_API_KEY", "test-key")
    routes = stock_route_names(libreoffice_available=False)
    router = ConversionRouter(routes=build_conversion_routes(route_names=routes))
    for mime, name in (
        (_DOCX, "office"),
        (_PPTX, "office"),
        ("application/pdf", "pdf"),
        ("message/rfc822", "email"),
        ("application/x-ipynb+json", "notebook"),
        ("text/html", "markitdown"),
        ("application/epub+zip", "markitdown"),
    ):
        assert router.converter_for(mime=mime).name == name, mime


def test_a_converter_that_overruns_its_time_limit_fails(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A parse past the wall-time limit fails the version with a typed error."""
    monkeypatch.setattr(time_limit, "CONVERTER_TIME_LIMIT_S", 0.05)
    with pytest.raises(ConversionError, match="conversion limit"):
        time_limit.run_with_time_limit(work=lambda: time.sleep(1), what="slow parse")
    assert time_limit.run_with_time_limit(work=lambda: 7, what="fast parse") == 7
    with pytest.raises(ValueError, match="boom"):
        time_limit.run_with_time_limit(work=_raise_boom, what="failing parse")


# --- helpers -----------------------------------------------------------------


class _FakeOcr:
    """A stand-in OCR route that counts its calls."""

    def __init__(
        self, *, metadata: DocumentMetadata | None = None, pages: tuple[int, ...] = (1,)
    ) -> None:
        self.calls = 0
        self._metadata = metadata
        self._pages = pages

    @property
    def name(self) -> str:
        return "fake_ocr"

    @property
    def version(self) -> str:
        return "fake-ocr-1"

    def convert(self, *, content: bytes, mime: str) -> ConversionResult:
        self.calls += 1
        return ConversionResult(
            document_md="OCR text\n",
            metadata=self._metadata,
            source_map=tuple(
                SourceMapEntry(
                    start=0, end=8, locators=(PageLocator(page=page, precision="page"),)
                )
                for page in self._pages
            ),
            manifest=ConverterManifest(
                components=(
                    ManifestComponent(
                        name="fake_ocr", version="1", execution="provider:fake"
                    ),
                ),
                coverage=ConversionCoverage(policy="ocr", complete=True),
                derivation_ranges=(),
            ),
        )


def _raise_boom() -> None:
    raise ValueError("boom")


def _assert_prose(*, result: ConversionResult) -> None:
    """The whole reading is labelled with eligible (prose) kinds."""
    assert result.manifest.derivation_ranges
    assert {
        labeled.derivation_kind for labeled in result.manifest.derivation_ranges
    } == {"prose"}


def _pptx() -> bytes:
    """A two-slide deck: title + body + notes, then a table; core properties."""
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    assert first.shapes.title is not None
    first.shapes.title.text = "Roadmap review"
    first.placeholders[1].text = "Launch in October."  # pyright: ignore[reportAttributeAccessIssue]
    notes = first.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Mention the budget."
    second = deck.slides.add_slide(deck.slide_layouts[6])
    table = second.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    for row, values in enumerate((("Region", "Revenue"), ("North", "1200"))):
        for column, value in enumerate(values):
            table.table.cell(row, column).text = value
    deck.core_properties.title = "Q4 roadmap"
    deck.core_properties.author = "Alice Novak"
    deck.core_properties.created = datetime(2026, 1, 2, 3, 4, 5)
    deck.core_properties.modified = datetime(2026, 2, 3, 4, 5, 6)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


_DOCX_CORE = (
    b'<?xml version="1.0" encoding="UTF-8"?>'
    b'<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/'
    b'metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" '
    b'xmlns:dcterms="http://purl.org/dc/terms/" '
    b'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    b"<dc:title>Plan</dc:title><dc:creator>Bob Smith</dc:creator>"
    b'<dcterms:created xsi:type="dcterms:W3CDTF">2025-05-06T07:08:09Z</dcterms:created>'
    b"</cp:coreProperties>"
)


def _docx(*, core_xml: bytes = _DOCX_CORE) -> bytes:
    """A minimal Word package with one paragraph and core properties."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        package.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/'
            'vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        package.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Type="http://schemas.'
            "openxmlformats.org/officeDocument/2006/relationships/officeDocument"
            '" Target="word/document.xml"/></Relationships>',
        )
        package.writestr(
            "word/document.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main"><w:body><w:p><w:r>'
            "<w:t>Ship the converter.</w:t></w:r></w:p></w:body></w:document>",
        )
        package.writestr("docProps/core.xml", core_xml)
    return buffer.getvalue()


def _with_replaced_member(
    *, content: bytes, member: str, old: bytes, new: bytes
) -> bytes:
    """Rewrite one zip member's bytes."""
    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == member:
                data = data.replace(old, new)
            target.writestr(info, data)
    return buffer.getvalue()


def _pdf(
    *,
    pages: list[str | None],
    info: dict[str, str] | None = None,
    to_unicode: dict[str, str] | None = None,
) -> bytes:
    """A small valid PDF: one Helvetica text line per page (None = no text).

    ``to_unicode`` maps hex byte codes to UTF-16BE hex for a ToUnicode CMap.
    """
    font = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica"
    objects: dict[int, str] = {1: "<< /Type /Catalog /Pages 2 0 R >>"}
    number = 4
    if to_unicode:
        pairs = "\n".join(f"<{code}> <{uni}>" for code, uni in to_unicode.items())
        cmap = (
            "/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            "/CMapName /Adobe-Identity-UCS def /CMapType 2 def\n"
            "1 begincodespacerange <00> <FF> endcodespacerange\n"
            f"{len(to_unicode)} beginbfchar\n{pairs}\nendbfchar\n"
            "endcmap CMapName currentdict /CMap defineresource pop end end"
        )
        objects[number] = f"<< /Length {len(cmap)} >>\nstream\n{cmap}\nendstream"
        font += f" /ToUnicode {number} 0 R"
        number += 1
    objects[3] = font + " >>"
    kids: list[str] = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else ""
        objects[number + 1] = (
            f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"
        )
        objects[number] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {number + 1} 0 R >>"
        )
        kids.append(f"{number} 0 R")
        number += 2
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    info_ref = ""
    if info:
        objects[number] = (
            "<< " + " ".join(f"/{key} ({value})" for key, value in info.items()) + " >>"
        )
        info_ref = f" /Info {number} 0 R"
    output = b"%PDF-1.4\n"
    offsets: dict[int, int] = {}
    for identifier in sorted(objects):
        offsets[identifier] = len(output)
        output += f"{identifier} 0 obj\n{objects[identifier]}\nendobj\n".encode(
            "latin-1"
        )
    xref = len(output)
    size = max(objects) + 1
    output += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    output += b"".join(
        f"{offsets[identifier]:010d} 00000 n \n".encode()
        for identifier in range(1, size)
    )
    output += (
        f"trailer\n<< /Size {size} /Root 1 0 R{info_ref} >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return output


def _epub() -> bytes:
    """A minimal EPUB with one XHTML chapter."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr("mimetype", "application/epub+zip")
        book.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        book.writestr(
            "content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
            'version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            "<dc:title>Tiny book</dc:title></metadata><manifest>"
            '<item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/>'
            '</manifest><spine><itemref idref="c1"/></spine></package>',
        )
        book.writestr(
            "c1.xhtml",
            '<html xmlns="http://www.w3.org/1999/xhtml"><body>'
            "<p>Chapter one begins.</p></body></html>",
        )
    return buffer.getvalue()


def _eml() -> bytes:
    """A multipart message with plain and HTML bodies and one attachment."""
    message = EmailMessage()
    message["From"] = "Alice Novak <alice@acme.com>"
    message["To"] = "Bob <bob@acme.com>"
    message["Cc"] = "carol@acme.com"
    message["Subject"] = "Budget approval"
    message["Date"] = "Wed, 04 Mar 2026 10:30:00 +0100"
    message["Message-ID"] = "<reply@acme.com>"
    message["References"] = "<root@acme.com> <middle@acme.com>"
    message.set_content("The budget is approved.\n")
    message.add_alternative("<p>The budget is <b>approved</b>.</p>", subtype="html")
    message.add_attachment(
        b"cells,data",
        maintype="application",
        subtype="octet-stream",
        filename="budget.xlsx",
    )
    return message.as_bytes()


def _ipynb() -> bytes:
    """A notebook: a Markdown cell, a code cell with output, a Markdown cell."""
    return json.dumps(
        {
            "nbformat": 4,
            "nbformat_minor": 5,
            "metadata": {"language_info": {"name": "python"}},
            "cells": [
                {"cell_type": "markdown", "source": ["# Churn study\n", "Setup."]},
                {
                    "cell_type": "code",
                    "source": "import pandas as pd",
                    "outputs": [{"output_type": "stream", "text": "SECRET OUTPUT"}],
                },
                {"cell_type": "markdown", "source": "Churn fell by 3%."},
                {"cell_type": "code", "source": "", "outputs": []},
            ],
        }
    ).encode()
